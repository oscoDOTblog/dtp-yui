"""Legacy plain-text resume builder (ReportLab / line dump)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .tailor import PROJECT_PRIORITY


def select_work_bullets(
    work_history: list[dict], role_family: str, limit: int = 8
) -> list[tuple[dict, list[str]]]:
    ordered = sorted(work_history, key=lambda w: w.get("startDate") or "", reverse=True)
    result = []
    for role in ordered:
        bullets = list(role.get("bullets") or [])
        result.append((role, bullets[: max(3, limit // max(len(ordered), 1))]))
    return result


def select_project_bullets(
    projects: list[dict], role_family: str
) -> list[tuple[dict, list[str]]]:
    priority = PROJECT_PRIORITY.get(
        role_family, PROJECT_PRIORITY["product"]
    )
    by_id = {p["_id"]: p for p in projects}
    selected = []
    for pid in priority:
        if pid in by_id:
            selected.append(by_id[pid])
    for p in projects:
        if p not in selected:
            selected.append(p)
        if len(selected) >= 4:
            break
    return [(p, list(p.get("resumeBullets") or [])[:3]) for p in selected]


def skills_line(skills: list[dict], categories: list[str], limit: int = 12) -> str:
    names = []
    for skill in skills:
        if skill.get("category") in categories and skill.get("approvedForResume"):
            names.append(skill["name"])
        if len(names) >= limit:
            break
    if len(names) < 6:
        for skill in skills:
            if skill.get("approvedForResume") and skill["name"] not in names:
                names.append(skill["name"])
            if len(names) >= limit:
                break
    return ", ".join(names)


def build_resume_lines(
    candidate: dict,
    work_history: list[dict],
    projects: list[dict],
    skills: list[dict],
    match: dict,
    job: dict,
) -> list[str]:
    role_family = match.get("roleFamily") or "product"
    positioning = (candidate.get("positioningSummaries") or {}).get(
        role_family
    ) or (candidate.get("positioningSummaries") or {}).get("product", "")

    contact = " | ".join(
        filter(
            None,
            [
                candidate.get("location"),
                candidate.get("email"),
                candidate.get("phone"),
                candidate.get("linkedin"),
                candidate.get("github"),
            ],
        )
    )

    lines = [
        candidate.get("name") or "Candidate",
        contact,
        "",
        "PROFESSIONAL SUMMARY",
        positioning,
        "",
        "TECHNICAL SKILLS",
    ]

    if role_family == "systems":
        lines.append(
            "Systems & Infrastructure: "
            + skills_line(skills, ["systems", "iac", "devops", "cicd", "cloud"], 14)
        )
        lines.append(
            "Data & Observability: "
            + skills_line(skills, ["data", "observability", "databases"], 10)
        )
        lines.append("Programming: " + skills_line(skills, ["languages", "backend"], 10))
    elif role_family == "mobile":
        lines.append("Mobile: " + skills_line(skills, ["mobile", "media"], 12))
        lines.append("Languages: " + skills_line(skills, ["languages"], 8))
        lines.append(
            "Backend/Cloud: " + skills_line(skills, ["cloud", "backend", "databases"], 10)
        )
    else:
        lines.append(
            "Frontend/Product: "
            + skills_line(skills, ["frontend", "product", "auth", "payments"], 12)
        )
        lines.append(
            "Backend/Cloud: "
            + skills_line(skills, ["backend", "cloud", "databases", "devops"], 12)
        )
        lines.append(
            "Languages: " + skills_line(skills, ["languages", "mobile", "ai"], 10)
        )

    lines += ["", "PROFESSIONAL EXPERIENCE"]
    for role, bullets in select_work_bullets(work_history, role_family):
        header = f"{role.get('company')} — {role.get('companyLocation') or ''}".strip(
            " —"
        )
        lines.append(header)
        lines.append(
            f"{role.get('title')} | {role.get('startDate')} – {role.get('endDate') or 'Present'}"
        )
        for b in bullets:
            lines.append(f"• {b}")
        lines.append("")

    lines.append("INDEPENDENT SYSTEMS AND PRODUCT ENGINEERING")
    lines.append(
        "Founder & Software Engineer | Oakland, California | November 2025 – Present"
    )
    for project, bullets in select_project_bullets(projects, role_family):
        lines.append(project.get("name") or "Project")
        for b in bullets:
            lines.append(f"• {b}")
        if not bullets and project.get("summary"):
            lines.append(f"• {project['summary']}")
        lines.append("")

    edu = (candidate.get("education") or [{}])[0]
    lines += [
        "EDUCATION",
        edu.get("institution") or "",
        f"{edu.get('degree') or ''} | {edu.get('school') or ''} | {edu.get('location') or ''} | {edu.get('graduatedAt') or ''}",
    ]
    return lines


def build_tailored_resume_lines(
    *,
    candidate: dict,
    skills: list[dict],
    work_history: list[dict],
    projects: list[dict],
    catalog: list,
    payload,
    job: dict | None = None,
) -> list[str]:
    """Plain-text resume lines from a verified TailorPayload (legacy PDF/txt)."""
    from .layout import build_resume_layout, layout_to_plain_lines

    layout = build_resume_layout(
        candidate=candidate,
        skills=skills,
        work_history=work_history,
        projects=projects,
        catalog=catalog,
        payload=payload,
    )
    return layout_to_plain_lines(layout)


def _format_skill_lines(payload, skill_by_id: dict) -> list[str]:
    """Grouped skill lines when payload.skillsGrouped is present."""
    grouped = getattr(payload, "skillsGrouped", None) or {}
    if grouped:
        lines: list[str] = []
        for cat, ids in grouped.items():
            names = [
                skill_by_id[sid]["name"]
                for sid in ids
                if sid in skill_by_id and skill_by_id[sid].get("name")
            ]
            if names:
                lines.append(f"{cat}: {', '.join(names)}")
        if lines:
            return lines
    skill_names = [
        skill_by_id[sid]["name"]
        for sid in payload.selectedSkillIds
        if sid in skill_by_id and skill_by_id[sid].get("name")
    ]
    if skill_names:
        return [", ".join(skill_names)]
    return []


def paragraphs_to_docx(path: Path, lines: list[str]) -> None:
    doc = Document()
    for line in lines:
        if not line.strip():
            doc.add_paragraph("")
        else:
            doc.add_paragraph(line)
    doc.save(str(path))


def paragraphs_to_pdf(path: Path, lines: list[str]) -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        spaceAfter=6,
    )
    header = ParagraphStyle(
        "Header",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        spaceAfter=10,
    )
    story = []
    for i, line in enumerate(lines):
        text = (line or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if not text.strip():
            story.append(Spacer(1, 0.12 * inch))
            continue
        style = header if i == 0 else body
        story.append(Paragraph(text, style))
    pdf = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )
    pdf.build(story)
