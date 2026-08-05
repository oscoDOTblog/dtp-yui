"""Structured DOCX from validated TailorPayload (not PDF→Word)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from .achievements import Achievement, catalog_by_id
from .tailor import TailorPayload


def render_tailored_docx(
    path: Path,
    *,
    candidate: dict,
    skills: list[dict],
    work_history: list[dict],
    projects: list[dict],
    catalog: list[Achievement],
    payload: TailorPayload,
    job: dict | None = None,
) -> None:
    by_id = catalog_by_id(catalog)
    rewrite = payload.rewrite_map()
    skill_by_id = {s["_id"]: s for s in skills if s.get("_id")}
    work_by_id = {w["_id"]: w for w in work_history if w.get("_id")}
    project_by_id = {p["_id"]: p for p in projects if p.get("_id")}

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    doc.add_heading(candidate.get("name") or "Candidate", level=0)
    contact = " | ".join(
        filter(
            None,
            [
                candidate.get("location"),
                candidate.get("email"),
                candidate.get("phone"),
                candidate.get("linkedin"),
                candidate.get("github"),
            ],
        )
    )
    if contact:
        doc.add_paragraph(contact)

    if payload.targetRole:
        p = doc.add_paragraph()
        run = p.add_run(payload.targetRole)
        run.bold = True

    if payload.summary:
        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph(payload.summary)

    highlights = list(getattr(payload, "highlights", None) or [])
    if highlights:
        doc.add_heading("Selected Highlights", level=1)
        for h in highlights:
            doc.add_paragraph(h.text, style="List Bullet")

    skill_groups = getattr(payload, "skillsGrouped", None) or {}
    if skill_groups:
        doc.add_heading("Technical Skills", level=1)
        for cat, ids in skill_groups.items():
            names = [
                skill_by_id[sid]["name"]
                for sid in ids
                if sid in skill_by_id and skill_by_id[sid].get("name")
            ]
            if names:
                doc.add_paragraph(f"{cat}: {', '.join(names)}")
    else:
        skill_names = [
            skill_by_id[sid]["name"]
            for sid in payload.selectedSkillIds
            if sid in skill_by_id and skill_by_id[sid].get("name")
        ]
        if skill_names:
            doc.add_heading("Technical Skills", level=1)
            doc.add_paragraph(", ".join(skill_names))

    # Group work achievements by parent role
    work_groups: dict[str, list[str]] = {}
    work_order: list[str] = []
    for aid in payload.selectedAchievementIds:
        ach = by_id.get(aid)
        if not ach or ach.kind != "work":
            continue
        if ach.parentId not in work_groups:
            work_groups[ach.parentId] = []
            work_order.append(ach.parentId)
        work_groups[ach.parentId].append(rewrite.get(aid) or ach.statement)

    if work_order:
        doc.add_heading("Professional Experience", level=1)
        for wid in work_order:
            role = work_by_id.get(wid) or {}
            header = f"{role.get('company') or 'Employer'} — {role.get('title') or ''}"
            dates = f"{role.get('startDate') or ''} – {role.get('endDate') or 'Present'}"
            p = doc.add_paragraph()
            run = p.add_run(header.strip(" —"))
            run.bold = True
            doc.add_paragraph(dates)
            for bullet in work_groups[wid]:
                doc.add_paragraph(bullet, style="List Bullet")

    project_groups: dict[str, list[str]] = {}
    project_order: list[str] = []
    for aid in payload.selectedAchievementIds:
        ach = by_id.get(aid)
        if not ach or ach.kind != "project":
            continue
        if ach.parentId not in project_groups:
            project_groups[ach.parentId] = []
            project_order.append(ach.parentId)
        project_groups[ach.parentId].append(rewrite.get(aid) or ach.statement)

    if project_order:
        doc.add_heading("Independent Projects", level=1)
        doc.add_paragraph(
            "Founder & Software Engineer | Oakland, California | November 2025 – Present"
        )
        for pid in project_order:
            project = project_by_id.get(pid) or {}
            name = project.get("name") or pid
            p = doc.add_paragraph()
            run = p.add_run(name)
            run.bold = True
            for bullet in project_groups[pid]:
                doc.add_paragraph(bullet, style="List Bullet")

    edu = (candidate.get("education") or [{}])[0]
    if edu:
        doc.add_heading("Education", level=1)
        if edu.get("institution"):
            doc.add_paragraph(str(edu.get("institution")))
        edu_line = " | ".join(
            filter(
                None,
                [
                    edu.get("degree"),
                    edu.get("school"),
                    edu.get("location"),
                    edu.get("graduatedAt"),
                ],
            )
        )
        if edu_line:
            doc.add_paragraph(edu_line)

    if job:
        doc.add_paragraph(
            f"Tailored for: {job.get('title')} at {job.get('company')}"
        )

    doc.save(str(path))
